import tempfile
import unittest
from datetime import datetime
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zoneinfo import ZoneInfo

import banco
from openpyxl import load_workbook
from relatorios.relatorio import criar_relatorio


class ControleHorasTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.db_original = banco.BANCO
        banco.BANCO = Path(self.temp_dir.name) / "teste.db"
        banco.criar_banco()

    def tearDown(self):
        banco.BANCO = self.db_original
        self.temp_dir.cleanup()

    def test_bloqueia_entrada_duplicada_aberta(self):
        with patch.object(banco, "agora", return_value=datetime(2026, 10, 19, 9, 0, tzinfo=ZoneInfo("America/Sao_Paulo"))):
            self.assertTrue(banco.registrar_entrada("Pessoa Teste"))
            self.assertFalse(banco.registrar_entrada("Pessoa Teste"))
        self.assertEqual(len(banco.listar_registros()), 1)

    def test_saida_sem_entrada_aberta_retorna_falso(self):
        self.assertFalse(banco.registrar_saida("Pessoa Teste"))

    def test_calcula_jornada_que_cruza_meia_noite_antes_de_oito_horas(self):
        with patch.object(banco, "agora", return_value=datetime(2026, 10, 18, 22, 0, tzinfo=ZoneInfo("America/Sao_Paulo"))):
            self.assertTrue(banco.registrar_entrada("Pessoa Teste"))
        with patch.object(banco, "agora", return_value=datetime(2026, 10, 19, 5, 0, tzinfo=ZoneInfo("America/Sao_Paulo"))):
            self.assertTrue(banco.registrar_saida("Pessoa Teste"))
        registro = banco.listar_registros()[0]
        self.assertEqual(registro["total"], "07:00")
        self.assertEqual(registro["saida"], "05:00")

    def test_encerramento_automatico_apos_oito_horas_fixa_total_em_duas(self):
        with patch.object(banco, "agora", return_value=datetime(2026, 10, 18, 22, 0, tzinfo=ZoneInfo("America/Sao_Paulo"))):
            self.assertTrue(banco.registrar_entrada("Pessoa Teste"))
        with patch.object(banco, "agora", return_value=datetime(2026, 10, 19, 6, 1, tzinfo=ZoneInfo("America/Sao_Paulo"))):
            self.assertEqual(banco.encerrar_jornadas_expiradas(), 1)
        registro = banco.listar_registros()[0]
        self.assertEqual(registro["saida"], "06:00")
        self.assertEqual(registro["total"], "02:00")
        self.assertEqual(banco.encerrar_jornadas_expiradas(), 0)

    def test_membro_adicionado_e_removido_sem_apagar_historico(self):
        ok, _ = banco.adicionar_membro("Pessoa Nova Teste")
        self.assertTrue(ok)
        self.assertIn("Pessoa Nova Teste", banco.listar_membros())
        ok, _ = banco.excluir_membro("Pessoa Nova Teste")
        self.assertTrue(ok)
        self.assertNotIn("Pessoa Nova Teste", banco.listar_membros())

    def test_relatorio_excel_e_gerado_em_memoria(self):
        with patch.object(banco, "agora", return_value=datetime(2026, 10, 19, 9, 0, tzinfo=ZoneInfo("America/Sao_Paulo"))):
            banco.registrar_entrada("Pessoa Teste")
        arquivo = criar_relatorio()
        self.assertIsInstance(arquivo, BytesIO)
        arquivo.seek(0)
        workbook = load_workbook(arquivo, read_only=True)
        self.assertEqual(workbook.sheetnames, ["Registros", "Resumo", "Ranking"])
        self.assertEqual(workbook["Registros"]["A1"].value, "Nome")
        workbook.close()


if __name__ == "__main__":
    unittest.main()
