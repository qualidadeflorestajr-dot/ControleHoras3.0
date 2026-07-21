from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime

from banco import (
    listar_registros,
    total_por_membro,
    total_geral
)


# ============================================
# CORES
# ============================================

VERDE = PatternFill(
    start_color="166534",
    end_color="166534",
    fill_type="solid"
)

BRANCO = Font(
    color="FFFFFF",
    bold=True
)


# ============================================
# FORMATAR CABEÇALHO
# ============================================

def formatar_cabecalho(planilha):

    for celula in planilha[1]:

        celula.fill = VERDE

        celula.font = BRANCO

        celula.alignment = Alignment(horizontal="center")


# ============================================
# CRIAR RELATÓRIO
# ============================================

def criar_relatorio():

    wb = Workbook()

    aba = wb.active

    aba.title = "Registros"

    aba.append([

        "Nome",
        "Data",
        "Entrada",
        "Saída",
        "Total"

    ])

    formatar_cabecalho(aba)

    registros = listar_registros()

    for registro in registros:

        aba.append([

            registro["nome"],
            registro["data"],
            registro["entrada"],
            registro["saida"],
            registro["total"]

        ])
        # ============================================
# ABA RESUMO
# ============================================

    resumo = wb.create_sheet("Resumo")

    resumo.append(["Resumo Semanal"])
    resumo["A1"].font = Font(bold=True, size=14)

    resumo.append([])
    resumo.append(["Total Geral", total_geral()])
    resumo.append(["Quantidade de Registros", len(registros)])
    resumo.append(["Quantidade de Membros", len(total_por_membro())])



# ============================================
# ABA RANKING
# ============================================

    ranking = wb.create_sheet("Ranking")

    ranking.append([

        "Posição",
        "Membro",
        "Horas"

    ])

    formatar_cabecalho(ranking)

    for posicao, membro in enumerate(total_por_membro(), start=1):

        ranking.append([

            posicao,
            membro["nome"],
            membro["total"]

        ])



# ============================================
# AJUSTAR LARGURA DAS COLUNAS
# ============================================

    for planilha in wb.worksheets:

        for coluna in planilha.columns:

            tamanho = 0

            letra = coluna[0].column_letter

            for celula in coluna:

                try:

                    if len(str(celula.value)) > tamanho:

                        tamanho = len(str(celula.value))

                except:

                    pass

            planilha.column_dimensions[letra].width = tamanho + 5
            # ============================================
# SALVAR O RELATÓRIO
# ============================================

    import os

    pasta = "relatorios"

    if not os.path.exists(pasta):

        os.makedirs(pasta)


    data_arquivo = datetime.now().strftime("%d-%m-%Y")

    caminho = os.path.join(

        pasta,

        f"Controle_Horas_{data_arquivo}.xlsx"

    )


    wb.save(caminho)

    return caminho