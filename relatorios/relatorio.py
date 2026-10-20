from io import BytesIO
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from banco import listar_registros, total_por_membro, total_geral
from utils import minutos_para_hora, total_semanal


VERDE = PatternFill(fill_type="solid", start_color="166534", end_color="166534")
BRANCO = Font(color="FFFFFF", bold=True)
FUSO_BRASIL = ZoneInfo("America/Sao_Paulo")


def formatar_cabecalho(planilha):
    for celula in planilha[1]:
        celula.fill = VERDE
        celula.font = BRANCO
        celula.alignment = Alignment(horizontal="center", vertical="center")
    planilha.freeze_panes = "A2"
    planilha.auto_filter.ref = planilha.dimensions


def ajustar_colunas(planilha):
    for coluna in planilha.columns:
        maior = max((len(str(c.value or "")) for c in coluna), default=0)
        letra = get_column_letter(coluna[0].column)
        planilha.column_dimensions[letra].width = min(max(maior + 3, 12), 42)


def criar_relatorio():
    """Gera o Excel em memória, sem depender de permissões de escrita no servidor."""
    wb = Workbook()
    registros = listar_registros()

    aba = wb.active
    aba.title = "Registros"
    aba.append(["Nome", "Data", "Entrada", "Saída", "Total", "Distância"])
    for registro in registros:
        aba.append([
            registro["nome"],
            registro["data"],
            registro["entrada"] or "",
            registro["saida"] or "",
            registro["total"] or "",
            registro["distancia"] or "Não informado",
        ])
    formatar_cabecalho(aba)

    resumo = wb.create_sheet("Resumo")
    resumo.append(["Indicador", "Valor"])
    resumo.append(["Total acumulado", total_geral()])
    resumo.append(["Total dos últimos 7 dias", minutos_para_hora(sum(item["minutos"] for item in total_semanal()))])
    resumo.append(["Registros", len(registros)])
    resumo.append(["Membros com horas registradas", len(total_por_membro())])
    resumo["B3"].number_format = "@"
    formatar_cabecalho(resumo)

    ranking = wb.create_sheet("Ranking")
    ranking.append(["Posição", "Membro", "Horas acumuladas"])
    for posicao, membro in enumerate(total_por_membro(), start=1):
        ranking.append([posicao, membro["nome"], membro["total"]])
    formatar_cabecalho(ranking)

    for planilha in wb.worksheets:
        ajustar_colunas(planilha)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def nome_arquivo_relatorio():
    data_arquivo = datetime.now(FUSO_BRASIL).strftime("%d-%m-%Y")
    return f"Controle_Horas_{data_arquivo}.xlsx"
